"""Label handling: video-id parsing, free-text timestamp parsing, labels.csv I/O.

Labels are in `open/open_labels.xlsx` (sheet `labels`), columns:
    A: file name ("X-1", "Х-3" -- Latin or Cyrillic first letter),
    B: timestamp as free text ("22 секунда", "1минута 8 секунда", may be empty),
    C: "нарушение" / "без нарушения".
Videos and labels are matched by the integer id only (see ASSUMPTIONS.md A-04).
"""

from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Optional

LABEL_VIOLATION = "violation"
LABEL_CLEAN = "no_violation"

_ID_RE = re.compile(r"(\d+)\s*$")
_VIOLATION_WORDS = ("наруш", "violation")
_CLEAN_WORDS = ("без наруш", "нет наруш", "no_violation", "no violation", "clean")

# Russian/English unit words -> seconds multiplier
_UNIT_PATTERNS = [
    (re.compile(r"^(?:ч|час|часа|часов|h|hr|hour|hours)$", re.I), 3600.0),
    (re.compile(r"^(?:м|мин|минута|минуты|минут|минуту|m|min|mins|minute|minutes)$", re.I), 60.0),
    (re.compile(r"^(?:с|сек|секунда|секунды|секунд|секунду|секунде|s|sec|secs|second|seconds)$", re.I), 1.0),
]


def parse_video_id(name: str) -> Optional[int]:
    """Return the integer id from names like 'X-1', 'Х-10.mov', 'public/Х-47.mov'.

    The leading letter may be Latin X (U+0058) or Cyrillic Х (U+0425); it is ignored.
    Returns None if no trailing integer is found.
    """
    stem = Path(str(name).strip()).name
    stem = re.sub(r"\.(mov|mp4|avi|mkv|m4v)$", "", stem, flags=re.I).strip()
    m = _ID_RE.search(stem)
    return int(m.group(1)) if m else None


def parse_timestamp(text) -> Optional[float]:
    """Parse a free-text timestamp into float seconds.

    Accepts: "22 секунда", "1минута 8 секунда", "1 мин 8 сек", "1:08", "01:02:03",
    "68", "68.5", "68,5", "11-я секунда", "на 22 секунде", None/"" -> None.
    Raises ValueError on text that contains no number.
    """
    if text is None:
        return None
    if isinstance(text, (int, float)):
        return float(text)
    s = str(text).strip().lower().replace(" ", " ")
    if not s:
        return None
    # "mm:ss" or "hh:mm:ss"
    if re.fullmatch(r"\d{1,2}(:\d{1,2}){1,2}(?:[.,]\d+)?", s):
        parts = s.replace(",", ".").split(":")
        total = 0.0
        for p in parts:
            total = total * 60.0 + float(p)
        return total
    # strip ordinal suffixes like "11-я", "11-ой", "22-й"
    s = re.sub(r"(\d)\s*-\s*(я|й|ой|ая|ую|ом|ем|е|и)\b", r"\1", s)
    # decimal comma -> dot (only when between digits)
    s = re.sub(r"(\d),(\d)", r"\1.\2", s)
    # tokens: number followed (optionally without space) by a unit word
    tokens = re.findall(r"(\d+(?:\.\d+)?)\s*([a-zа-яё]*)", s)
    if not tokens:
        raise ValueError(f"no number in timestamp text: {text!r}")
    total = 0.0
    any_unit = False
    bare: list[float] = []
    for num, unit in tokens:
        val = float(num)
        mult = None
        if unit:
            for pat, m in _UNIT_PATTERNS:
                if pat.match(unit):
                    mult = m
                    break
        if mult is None:
            bare.append(val)
        else:
            any_unit = True
            total += val * mult
    if not any_unit:
        # plain number(s): "68" -> 68 s
        return bare[0]
    # a bare number that follows a minutes unit is seconds ("1 мин 8")
    if bare and len(tokens) >= 2:
        total += bare[-1]
    return total


def normalize_label(text) -> str:
    """Map the Russian label text to LABEL_VIOLATION / LABEL_CLEAN."""
    s = str(text or "").strip().lower()
    for w in _CLEAN_WORDS:
        if w in s:
            return LABEL_CLEAN
    for w in _VIOLATION_WORDS:
        if w in s:
            return LABEL_VIOLATION
    raise ValueError(f"unrecognized label text: {text!r}")


def file_md5(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        while True:
            b = f.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


@dataclass
class LabelRow:
    video_id: int
    path: str  # relative path to the video file, '' if the file is missing
    label: str  # LABEL_VIOLATION | LABEL_CLEAN | '' for unlabeled
    t_sec: Optional[float]  # labeled event time, None for clean videos
    split: str  # 'labeled' | 'unlabeled'
    content_md5: str  # md5 of the file ('' if missing); groups duplicates for LOVO
    raw_name: str = ""
    raw_timestamp: str = ""
    raw_label: str = ""


def read_labels_xlsx(xlsx_path: Path, sheet: str = "labels") -> list[dict]:
    """Read the raw label rows from the workbook (requires openpyxl)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # header
        if row is None or all(c in (None, "") for c in row):
            continue
        name, ts, lab = (list(row) + [None, None, None])[:3]
        if name in (None, ""):
            continue
        rows.append(
            {
                "raw_name": str(name),
                "raw_timestamp": "" if ts is None else str(ts),
                "raw_label": str(lab or ""),
            }
        )
    return rows


def index_videos(video_dir: Path, exts: Iterable[str] = (".mov", ".mp4")) -> dict[int, Path]:
    """Map integer id -> video path for every video file in `video_dir`."""
    out: dict[int, Path] = {}
    allowed = {e.lower() for e in exts}
    for p in sorted(Path(video_dir).iterdir()):
        if p.suffix.lower() not in allowed:
            continue
        vid = parse_video_id(p.name)
        if vid is None:
            continue
        if vid in out:
            raise ValueError(f"duplicate video id {vid}: {out[vid]} and {p}")
        out[vid] = p
    return out


def build_label_rows(
    raw_rows: list[dict], videos: dict[int, Path], repo_root: Path, with_md5: bool = True
) -> list[LabelRow]:
    """Join raw label rows with video files; add unlabeled videos as a separate split."""
    rows: list[LabelRow] = []
    seen: set[int] = set()
    md5_cache: dict[Path, str] = {}

    def md5_of(p: Optional[Path]) -> str:
        if p is None or not with_md5:
            return ""
        if p not in md5_cache:
            md5_cache[p] = file_md5(p)
        return md5_cache[p]

    def rel(p: Path) -> str:
        try:
            return p.resolve().relative_to(Path(repo_root).resolve()).as_posix()
        except ValueError:
            return p.as_posix()

    for r in raw_rows:
        vid = parse_video_id(r["raw_name"])
        if vid is None:
            raise ValueError(f"cannot parse video id from label row {r!r}")
        if vid in seen:
            raise ValueError(f"duplicate label row for id {vid}")
        seen.add(vid)
        label = normalize_label(r["raw_label"])
        t = parse_timestamp(r["raw_timestamp"]) if r["raw_timestamp"] else None
        if label == LABEL_VIOLATION and t is None:
            raise ValueError(f"violation row without timestamp: {r!r}")
        if label == LABEL_CLEAN:
            t = None
        p = videos.get(vid)
        rows.append(
            LabelRow(
                vid, rel(p) if p else "", label, t, "labeled", md5_of(p), r["raw_name"], r["raw_timestamp"], r["raw_label"]
            )
        )
    for vid, p in sorted(videos.items()):
        if vid not in seen:
            rows.append(LabelRow(vid, rel(p), "", None, "unlabeled", md5_of(p), p.name, "", ""))
    rows.sort(key=lambda r: r.video_id)
    return rows


CSV_FIELDS = ["video_id", "path", "label", "t_sec", "split", "content_md5", "raw_name", "raw_timestamp", "raw_label"]


def write_labels_csv(rows: list[LabelRow], out_path: Path) -> None:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        w.writeheader()
        for r in rows:
            d = asdict(r)
            d["t_sec"] = "" if r.t_sec is None else f"{r.t_sec:g}"
            w.writerow(d)


def read_labels_csv(path: Path) -> list[LabelRow]:
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        for d in csv.DictReader(f):
            rows.append(
                LabelRow(
                    video_id=int(d["video_id"]),
                    path=d["path"],
                    label=d["label"],
                    t_sec=float(d["t_sec"]) if d.get("t_sec") not in (None, "") else None,
                    split=d["split"],
                    content_md5=d.get("content_md5", ""),
                    raw_name=d.get("raw_name", ""),
                    raw_timestamp=d.get("raw_timestamp", ""),
                    raw_label=d.get("raw_label", ""),
                )
            )
    return rows


def lovo_groups(rows: list[LabelRow]) -> dict[str, list[int]]:
    """Leave-one-video-out groups keyed by content hash (duplicates share a fold)."""
    groups: dict[str, list[int]] = {}
    for r in rows:
        if r.split != "labeled":
            continue
        key = r.content_md5 or f"id-{r.video_id}"
        groups.setdefault(key, []).append(r.video_id)
    return groups
